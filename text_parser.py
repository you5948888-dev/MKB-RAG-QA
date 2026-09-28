"""
文本解析工具。
将上传的文档（PDF / TXT / MD / DOCX）解析为纯文本。

分块逻辑（Advanced RAG 第 1 期起）已统一迁移到 `chunking` 模块，
本文件保留 load_text 与两个兼容入口，避免旧调用方一次性改动。
"""


def load_text(file_path: str) -> str:
    """根据文件扩展名解析文档，返回纯文本。"""
    if file_path.lower().endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    if file_path.lower().endswith(".docx"):
        # python-docx：遍历所有段落（含表格内文本）拼成纯文本
        from docx import Document
        doc = Document(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)

    # TXT / MD 直接读取
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def split_text(text: str, chunk_size: int = 800, chunk_overlap: int = 100) -> list[str]:
    """【兼容入口】等价于通用分块的默认配置，实际逻辑见 chunking 模块。"""
    import chunking

    items = chunking.chunk_text(text, chunk_size=chunk_size, chunk_overlap=chunk_overlap)
    return [item["content"] for item in items]


def parse_and_split(file_path: str, chunk_size: int = 800, chunk_overlap: int = 100) -> list[dict]:
    """【兼容入口】解析文档并按通用分块默认配置切分，实际逻辑见 chunking 模块。

    每个块包含 chunk_index、content、length、chars 四个字段。
    """
    import chunking

    return chunking.chunk_document(file_path, chunk_size=chunk_size, chunk_overlap=chunk_overlap)